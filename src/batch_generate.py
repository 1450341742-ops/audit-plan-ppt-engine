from __future__ import annotations
import argparse
import re
from pathlib import Path

from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.util import Inches

from parser import parse_excel
import renderer
from ai_summary import generate_ai_top5

BASE_DIR = Path(__file__).resolve().parent.parent
DEFAULT_OUTPUT = BASE_DIR / "output"
DEFAULT_TEMPLATE = BASE_DIR / "assets" / "template.pptx"
DEFAULT_OUTPUT.mkdir(exist_ok=True, parents=True)

HEADER_BLUE = RGBColor(91, 155, 213)
ROW_LIGHT_BLUE = RGBColor(222, 230, 242)
ROW_LIGHT_YELLOW = RGBColor(255, 242, 204)
FIXED_COVER_COMPANY = "北京xxxx医药科技有限公司"


def _set_cell_font_color(cell, color: RGBColor):
    try:
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.color.rgb = color
    except Exception:
        pass


def safe_stem(name: str) -> str:
    return re.sub(r'[<>:"/\\|?*]+', '_', name).strip().strip('.') or 'output'


def _patched_render_cover(slide, context):
    renderer._remove_text_shapes(slide)
    meta = context.get("meta", {})
    project = renderer._clean(meta.get("project_name", "—"))
    center = renderer._clean(meta.get("center_name", "—"))
    center_no = renderer._clean(meta.get("center_no", ""))
    audit_date = renderer._clean(meta.get("audit_date", "—"))
    title = f"{project}-{center}" + (f"（中心编号{center_no}）" if center_no else "")
    renderer._add_textbox(slide, 0.25, 3.58, 12.70, 0.82, title, 22, True, renderer.WHITE, PP_ALIGN.LEFT)
    renderer._add_textbox(slide, 0.25, 4.62, 12.70, 0.40, "中心稽查末次会议", 22, True, renderer.YELLOW, PP_ALIGN.LEFT)
    renderer._add_textbox(slide, 0.35, 5.33, 12.30, 0.28, f"时间：{audit_date}", 14, True, renderer.WHITE, PP_ALIGN.LEFT)
    renderer._add_textbox(slide, 0.35, 5.68, 12.30, 0.28, FIXED_COVER_COMPANY, 14, True, renderer.WHITE, PP_ALIGN.LEFT)


renderer._render_cover = _patched_render_cover


def _patched_extract_top5_risks(context: dict) -> list[dict]:
    ai_rows = generate_ai_top5(context)
    if ai_rows:
        return ai_rows[:5]
    return renderer._extract_top5_risks(context)


def _patched_render_risk_summary(slide, context):
    renderer._clear_issue_content(slide)
    risks = _patched_extract_top5_risks(context)
    renderer._add_textbox(slide, 0.45, 0.28, 12.40, 0.42, "TOP5高风险问题及核查应对建议", 22, True, renderer.BLACK, PP_ALIGN.LEFT)
    if not risks:
        renderer._add_textbox(slide, 0.75, 1.35, 11.80, 0.80, "本次上传文件中未识别到可用于提炼TOP5的问题内容，请复核Excel问题分类、问题描述和依据列是否完整。", 16, False, renderer.BLACK, PP_ALIGN.LEFT)
        return
    shape = slide.shapes.add_table(6, 4, Inches(0.45), Inches(0.90), Inches(12.45), Inches(6.10))
    tbl = shape.table
    widths = [0.60, 2.75, 4.15, 4.95]
    for i, w in enumerate(widths):
        tbl.columns[i].width = Inches(w)
    tbl.rows[0].height = Inches(0.40)
    for r in range(1, 6):
        tbl.rows[r].height = Inches(1.14)
    headers = ["排名", "高风险问题", "风险维度分析", "核查应对建议"]
    for c, h in enumerate(headers):
        renderer._set_cell(tbl.cell(0, c), h, 11, True, PP_ALIGN.CENTER)
        renderer._set_cell_fill(tbl.cell(0, c), HEADER_BLUE)
        _set_cell_font_color(tbl.cell(0, c), renderer.WHITE)
    for r in range(1, 6):
        if r <= len(risks):
            item = risks[r - 1]
            values = [str(r), item.get("risk", "—"), item.get("analysis", "—"), item.get("advice", "—")]
        else:
            values = [str(r), "—", "—", "—"]
        for c, v in enumerate(values):
            size = 7 if c in (2, 3) else 8
            align = PP_ALIGN.CENTER if c == 0 else PP_ALIGN.LEFT
            renderer._set_cell(tbl.cell(r, c), v, size, c == 0, align)
            renderer._set_cell_fill(tbl.cell(r, c), ROW_LIGHT_YELLOW if c == 0 else ROW_LIGHT_BLUE)


renderer._render_risk_summary = _patched_render_risk_summary


def render_one(excel_path: str | Path, output_dir: str | Path | None = None, template_path: str | Path | None = None) -> Path:
    excel_path = Path(excel_path)
    output_dir = Path(output_dir or DEFAULT_OUTPUT)
    output_dir.mkdir(parents=True, exist_ok=True)
    context = parse_excel(excel_path)
    out = output_dir / f"{safe_stem(excel_path.stem)}-V8.12无来源标识版.pptx"
    renderer.render_ppt(context, out, template_path=Path(template_path or DEFAULT_TEMPLATE))
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--excel", type=str)
    ap.add_argument("--excel_dir", type=str)
    ap.add_argument("--output_dir", type=str, default=str(DEFAULT_OUTPUT))
    ap.add_argument("--template", type=str, default=str(DEFAULT_TEMPLATE))
    args = ap.parse_args()
    files = []
    if args.excel:
        files.append(Path(args.excel))
    if args.excel_dir:
        d = Path(args.excel_dir)
        for ext in ("*.xlsx", "*.xlsm", "*.xls"):
            files += sorted(d.glob(ext))
    if not files:
        raise SystemExit("请提供 --excel 或 --excel_dir")
    for f in files:
        out = render_one(f, args.output_dir, template_path=args.template)
        print(out)


if __name__ == "__main__":
    main()
